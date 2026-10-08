"""Xem trước file Office trong Kho tư liệu (nhúng vào <iframe> ở FE).

GET /kb/sources/{source_id}/preview/{path} — `path` cùng quy ước với raw_file (đường dẫn tương đối trong raw_dir):
- Word / PowerPoint (.doc .docx .odt .rtf .ppt .pptx .odp) -> PDF bằng LibreOffice headless, trả inline.
- Bảng tính (.xls .xlsx .xlsm .ods .csv) -> một trang HTML tự chứa: mỗi sheet một tab (CSS thuần, không JS),
  lưới có tiêu đề cột A,B,C + số dòng, header dính, giữ ô gộp, giá trị đã tính, cắt ở MAX_ROWS × MAX_COLS.
  .xls / .ods được LibreOffice chuyển sang .xlsx trước.

Bộ nhớ đệm nằm NGOÀI raw_dir (API chi tiết nguồn liệt kê raw_dir bằng rglob): <data>/preview/<source_id>/,
khoá theo đường dẫn + mtime + kích thước file gốc. Xoá nguồn thì xoá luôn (pipeline.delete_source_data -> drop_cache).
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import os
import queue
import shutil
import subprocess
import tempfile
import threading
from datetime import date, datetime, time
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse

from ..auth import current_user
from ..config import PREVIEW_MAX_COLS, PREVIEW_MAX_ROWS, PREVIEW_TIMEOUT, preview_dir, soffice_bin_env
from . import pipeline
from .pipeline import raw_dir, sources

router = APIRouter(prefix="/api")

PDF_EXTS = {".doc", ".docx", ".odt", ".rtf", ".ppt", ".pptx", ".odp"}
SHEET_EXTS = {".xls", ".xlsx", ".xlsm", ".ods", ".csv"}
MAX_ROWS = PREVIEW_MAX_ROWS
MAX_COLS = PREVIEW_MAX_COLS
TIMEOUT = PREVIEW_TIMEOUT       # giây cho một lượt LibreOffice
RENDER_VERSION = "1"                                      # đổi khi đổi cách dựng HTML -> bỏ đệm cũ
CONCURRENCY = 2                                           # số LibreOffice chạy cùng lúc

_slots: queue.Queue[int] = queue.Queue()
for _i in range(CONCURRENCY):
    _slots.put(_i)
_locks: dict[str, threading.Lock] = {}
_locks_guard = threading.Lock()


def preview_root() -> Path:
    """Thư mục anh em của RAW_DIR (vd data/preview), đọc lúc gọi để test đổi RAW_DIR được."""
    env = preview_dir()
    return Path(env) if env else Path(pipeline.RAW_DIR).parent / "preview"


def drop_cache(source_id) -> None:
    shutil.rmtree(preview_root() / str(source_id), ignore_errors=True)


def soffice_bin() -> str | None:
    env = soffice_bin_env()
    if env:
        return env if Path(env).is_file() else None
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac if Path(mac).is_file() else None


def _lock_for(key: str) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(key, threading.Lock())


# ---------------------------------------------------------------- LibreOffice

def convert(src: Path, fmt: str, workdir: Path) -> Path:
    """Chuyển `src` sang `fmt` (pdf / xlsx) vào `workdir`, trả đường dẫn file kết quả."""
    binary = soffice_bin()
    if not binary:
        raise HTTPException(503, "Máy chủ chưa cài LibreOffice nên chưa xem trước được file này")
    # tên file tạm thuần ASCII: LibreOffice đặt tên kết quả theo tên file vào
    inp = workdir / f"input{src.suffix.lower()}"
    shutil.copyfile(src, inp)
    outdir = workdir / "out"
    outdir.mkdir()
    slot = _slots.get()
    try:
        # mỗi lượt đồng thời một profile riêng (giữ lại giữa các lần để khỏi khởi tạo lại) -> không khoá nhau
        profile = preview_root() / ".lo-profile" / str(slot)
        profile.mkdir(parents=True, exist_ok=True)
        cmd = [binary, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless", "--norestore",
               "--nologo", "--nodefault", "--nolockcheck", "--convert-to", fmt, "--outdir", str(outdir), str(inp)]
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=TIMEOUT, cwd=workdir)
        except subprocess.TimeoutExpired:
            raise HTTPException(504, f"Chuyển đổi file quá {TIMEOUT} giây, thử lại sau") from None
    finally:
        _slots.put(slot)
    out = outdir / f"input.{fmt}"
    if proc.returncode != 0 or not out.is_file() or out.stat().st_size == 0:
        raise HTTPException(422, "Không chuyển được file để xem trước (file hỏng hoặc định dạng không hỗ trợ)")
    return out


# ---------------------------------------------------------------- bảng tính -> HTML

def col_letter(i: int) -> str:
    """1 -> A, 27 -> AA."""
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def fmt_value(v, number_format: str = "General") -> tuple[str, bool]:
    """(chuỗi hiển thị, là số?)"""
    if v is None:
        return "", False
    if isinstance(v, bool):
        return ("TRUE" if v else "FALSE"), False
    if isinstance(v, datetime):
        if v.time() == time(0):
            return v.strftime("%d/%m/%Y"), True
        return v.strftime("%d/%m/%Y %H:%M" if not v.second else "%d/%m/%Y %H:%M:%S"), True
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y"), True
    if isinstance(v, time):
        return v.strftime("%H:%M" if not v.second else "%H:%M:%S"), True
    if isinstance(v, (int, float)):
        nf = (number_format or "General").split(";")[0]
        group = "," in nf and "#" in nf
        if "%" in nf:
            dec = _decimals(nf)
            return f"{v * 100:.{dec}f}%", True
        if isinstance(v, float) and not v.is_integer():
            dec = _decimals(nf) if nf != "General" and "0" in nf else None
            if dec is not None:
                return (f"{v:,.{dec}f}" if group else f"{v:.{dec}f}"), True
            s = f"{v:.10g}"
            if "e" in s:
                return s, True
            if group:
                whole, _, frac = s.partition(".")
                s = f"{int(whole):,}" + ("." + frac if frac else "")
            return s, True
        n = int(v)
        dec = _decimals(nf) if nf != "General" else 0
        if dec:
            return (f"{n:,.{dec}f}" if group else f"{n:.{dec}f}"), True
        return (f"{n:,}" if group else str(n)), True
    return str(v), False


def _decimals(nf: str) -> int:
    if "." not in nf:
        return 0
    frac = nf.split(".", 1)[1]
    n = 0
    for ch in frac:
        if ch in "0#":
            n += 1
        else:
            break
    return n


class Sheet:
    def __init__(self, name: str):
        self.name = name
        self.rows: list[list[tuple[str, bool] | None]] = []  # None = ô bị gộp vào ô khác
        self.spans: dict[tuple[int, int], tuple[int, int]] = {}  # (r, c) 0-based -> (rowspan, colspan)
        self.widths: dict[int, float] = {}
        self.total_rows = 0
        self.total_cols = 0
        self.hidden = False


def read_xlsx(path: Path) -> list[Sheet]:
    import openpyxl
    from openpyxl.utils import column_index_from_string
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001 — file hỏng / có mật khẩu
        raise HTTPException(422, f"Không đọc được bảng tính: {e.__class__.__name__}") from None
    out = []
    for ws in wb.worksheets:
        sh = Sheet(ws.title)
        sh.hidden = ws.sheet_state != "visible"
        sh.total_rows, sh.total_cols = ws.max_row or 0, ws.max_column or 0
        nrows, ncols = min(sh.total_rows, MAX_ROWS), min(sh.total_cols, MAX_COLS)
        grid: list[list] = []
        last_r = last_c = 0
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=nrows, max_col=ncols), 1):
            cells = []
            for c, cell in enumerate(row, 1):
                val = fmt_value(cell.value, cell.number_format)
                if val[0]:
                    last_r, last_c = max(last_r, r), max(last_c, c)
                cells.append(val)
            grid.append(cells)
        covered: set[tuple[int, int]] = set()
        for rng in ws.merged_cells.ranges:
            r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
            if r1 > nrows or c1 > ncols:
                continue
            r2, c2 = min(r2, nrows), min(c2, ncols)
            if grid[r1 - 1][c1 - 1][0]:   # ô gộp có chữ -> giữ đủ vùng gộp
                last_r, last_c = max(last_r, r2), max(last_c, c2)
            sh.spans[(r1 - 1, c1 - 1)] = (r2 - r1 + 1, c2 - c1 + 1)
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    if (rr, cc) != (r1, c1):
                        covered.add((rr - 1, cc - 1))
        # bỏ dòng / cột trống ở đuôi (định dạng thừa làm max_row phình)
        if sh.total_rows <= MAX_ROWS:
            sh.total_rows = last_r
        if sh.total_cols <= MAX_COLS:
            sh.total_cols = last_c
        grid = [row[:last_c] for row in grid[:last_r]]
        for (r, c) in covered:
            if r < len(grid) and c < len(grid[r]):
                grid[r][c] = None
        sh.spans = {k: (min(rs, last_r - k[0]), min(cs, last_c - k[1])) for k, (rs, cs) in sh.spans.items()
                    if k[0] < last_r and k[1] < last_c}
        sh.rows = grid
        for key, dim in ws.column_dimensions.items():
            if not dim.width:
                continue
            try:
                lo = dim.min or column_index_from_string(key)
            except ValueError:
                continue
            for i in range(lo, min(dim.max or lo, last_c) + 1):
                sh.widths[i] = dim.width
        out.append(sh)
    wb.close()
    return out


def _is_number(v: str) -> bool:
    s = v.strip().replace(",", "").replace(" ", "")
    if not s:
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


def read_csv(path: Path) -> list[Sheet]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    sh = Sheet(path.stem)
    csv.field_size_limit(10 * 1024 * 1024)
    rows = list(csv.reader(io.StringIO(text), dialect))
    sh.total_rows = len(rows)
    sh.total_cols = max((len(r) for r in rows), default=0)
    ncols = min(sh.total_cols, MAX_COLS)
    for r in rows[:MAX_ROWS]:
        cells = []
        for v in r[:ncols]:
            cells.append((v, _is_number(v)))
        sh.rows.append(cells)
    return [sh]


CSS = """
:root{--bg:#fff;--fg:#1f2328;--grid:#d0d7de;--head:#f3f4f6;--headfg:#57606a;--tab:#e9ecef;--tabon:#fff;
--accent:#1a7f37;--note:#fff8c5;--notefg:#6b5100}
@media (prefers-color-scheme:dark){:root{--bg:#0d1117;--fg:#e6edf3;--grid:#30363d;--head:#161b22;--headfg:#8b949e;
--tab:#21262d;--tabon:#0d1117;--accent:#3fb950;--note:#3b2e00;--notefg:#f2cc60}}
*{box-sizing:border-box}html,body{margin:0;height:100%;background:var(--bg);color:var(--fg);
font:13px/1.35 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif}
body{display:flex;flex-direction:column}
input.t{position:absolute;opacity:0;pointer-events:none}
.tabs{order:2;display:flex;gap:2px;overflow-x:auto;background:var(--head);border-top:1px solid var(--grid);
padding:0 4px;flex:none}
.tabs label{padding:6px 12px;cursor:pointer;white-space:nowrap;background:var(--tab);color:var(--headfg);
border:1px solid var(--grid);border-top:none;border-radius:0 0 6px 6px;margin-bottom:4px}
.tabs label.h{font-style:italic;opacity:.7}
.panes{order:1;flex:1;min-height:0;position:relative}
.pane{display:none;position:absolute;inset:0;flex-direction:column}
.note{flex:none;padding:6px 10px;background:var(--note);color:var(--notefg);border-bottom:1px solid var(--grid)}
.scroll{flex:1;overflow:auto}
table{border-collapse:separate;border-spacing:0;table-layout:fixed}
td,th{border-right:1px solid var(--grid);border-bottom:1px solid var(--grid);padding:2px 6px;overflow:hidden;
text-overflow:ellipsis;white-space:pre-wrap;word-break:break-word;vertical-align:top;max-width:480px}
thead th{position:sticky;top:0;z-index:2;background:var(--head);color:var(--headfg);font-weight:500;
text-align:center;white-space:nowrap}
tbody th{position:sticky;left:0;z-index:1;background:var(--head);color:var(--headfg);font-weight:400;
text-align:right;min-width:44px}
thead th.c{left:0;z-index:3;position:sticky}
td.n{text-align:right;white-space:nowrap}
.empty{padding:24px;color:var(--headfg)}
"""


def render_html(sheets: list[Sheet], title: str) -> str:
    e = html.escape
    rules, inputs, labels, panes = [], [], [], []
    first = next((i for i, s in enumerate(sheets) if not s.hidden), 0)
    for i, sh in enumerate(sheets):
        inputs.append(f'<input class="t" type="radio" name="tab" id="t{i}"{" checked" if i == first else ""}>')
        labels.append(f'<label for="t{i}"{" class=h" if sh.hidden else ""} title="{e(sh.name)}">{e(sh.name)}</label>')
        rules.append(f"#t{i}:checked~.panes #p{i}{{display:flex}}#t{i}:checked~.tabs label[for=t{i}]"
                     f"{{background:var(--tabon);color:var(--accent);font-weight:600}}")
        ncols = max((len(r) for r in sh.rows), default=0)
        notes = []
        if sh.total_rows > len(sh.rows):
            notes.append(f"{len(sh.rows):,}/{sh.total_rows:,} dòng")
        if sh.total_cols > ncols and sh.total_cols > MAX_COLS:
            notes.append(f"{ncols}/{sh.total_cols} cột")
        note = (f'<div class="note">Bảng quá lớn — chỉ hiển thị {" và ".join(notes)} đầu tiên. '
                f'Tải file gốc để xem đầy đủ.</div>') if notes else ""
        if not sh.rows:
            body = '<div class="empty">Sheet trống</div>'
        else:
            parts = ['<table><colgroup><col style="width:48px">']
            for c in range(1, ncols + 1):
                w = sh.widths.get(c)
                px = int(w * 7 + 5) if w else 90
                parts.append(f'<col style="width:{max(24, min(px, 480))}px">')
            parts.append('</colgroup><thead><tr><th class="c"></th>')
            parts.extend(f"<th>{col_letter(c)}</th>" for c in range(1, ncols + 1))
            parts.append("</tr></thead><tbody>")
            for r, row in enumerate(sh.rows):
                parts.append(f"<tr><th>{r + 1}</th>")
                for c in range(ncols):
                    cell = row[c] if c < len(row) else ("", False)
                    if cell is None:
                        continue
                    attrs = ' class="n"' if cell[1] else ""
                    span = sh.spans.get((r, c))
                    if span:
                        if span[0] > 1:
                            attrs += f' rowspan="{span[0]}"'
                        if span[1] > 1:
                            attrs += f' colspan="{span[1]}"'
                    parts.append(f"<td{attrs}>{e(cell[0])}</td>")
                parts.append("</tr>")
            parts.append("</tbody></table>")
            body = "".join(parts)
        panes.append(f'<div class="pane" id="p{i}">{note}<div class="scroll">{body}</div></div>')
    tabs_style = "" if len(sheets) > 1 else ".tabs{display:none}"
    return ("<!doctype html><html lang=\"vi\"><head><meta charset=\"utf-8\">"
            "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">"
            "<meta name=\"color-scheme\" content=\"light dark\">"
            f"<title>{e(title)}</title><style>{CSS}{''.join(rules)}{tabs_style}</style></head><body>"
            f"{''.join(inputs)}<div class=\"tabs\">{''.join(labels)}</div>"
            f"<div class=\"panes\">{''.join(panes)}</div></body></html>")


# ---------------------------------------------------------------- endpoint

def build(target: Path, ext: str, out: Path) -> None:
    """Dựng bản xem trước của `target` vào `out` (ghi nguyên tử)."""
    with tempfile.TemporaryDirectory(prefix="preview-") as tmp:
        work = Path(tmp)
        if ext in PDF_EXTS:
            result = convert(target, "pdf", work)
            data = result.read_bytes()
        else:
            if ext == ".csv":
                sheets = read_csv(target)
            elif ext in (".xls", ".ods"):
                sheets = read_xlsx(convert(target, "xlsx", work))
            else:
                sheets = read_xlsx(target)
            data = render_html(sheets, target.name).encode("utf-8")
        part = out.with_name(out.name + ".part")
        part.write_bytes(data)
        os.replace(part, out)


def cached_preview(source_id, rel: str, target: Path) -> Path:
    ext = target.suffix.lower()
    st = target.stat()
    path_key = hashlib.sha1(rel.encode("utf-8")).hexdigest()[:16]
    ver = hashlib.sha1(f"{st.st_mtime_ns}|{st.st_size}|{RENDER_VERSION}".encode()).hexdigest()[:12]
    folder = preview_root() / str(source_id)
    out = folder / f"{path_key}-{ver}{'.pdf' if ext in PDF_EXTS else '.html'}"
    if out.is_file():
        return out
    with _lock_for(str(out)):
        if out.is_file():   # request khác vừa dựng xong
            return out
        folder.mkdir(parents=True, exist_ok=True)
        build(target, ext, out)
        for old in folder.glob(f"{path_key}-*"):   # bản cũ của cùng file (file gốc đã đổi)
            if old != out:
                old.unlink(missing_ok=True)
    return out


@router.get("/kb/sources/{source_id}/preview/{path:path}")
def preview_file(source_id: str, path: str, user: dict = Depends(current_user)):
    from .routes import load   # nhập muộn: routes nhập nhiều module nặng
    src, _ = load(sources, source_id, user, "source.read", "nguồn")
    base = raw_dir(src["_id"]).resolve()
    target = (base / path).resolve()
    if base not in target.parents or not target.is_file():
        raise HTTPException(404, "Không tìm thấy file")
    ext = target.suffix.lower()
    if ext not in PDF_EXTS | SHEET_EXTS:
        raise HTTPException(415, "Chưa hỗ trợ xem trước loại file này (chỉ Word, PowerPoint, Excel, CSV)")
    rel = str(target.relative_to(base))
    out = cached_preview(src["_id"], rel, target)
    headers = {"Cache-Control": "private, no-cache", "X-Content-Type-Options": "nosniff"}
    if ext in PDF_EXTS:
        name = target.stem + ".pdf"
        headers["Content-Disposition"] = f"inline; filename*=UTF-8''{quote(name)}"
        return FileResponse(out, media_type="application/pdf", headers=headers)
    # trang tự chứa, không script -> chặn mọi tài nguyên ngoài
    headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'"
    return FileResponse(out, media_type="text/html; charset=utf-8", headers=headers)
