"""Word, PowerPoint, Excel -> Markdown, không cần AI. Dùng chung cho file Office tải lên và Google Docs / Sheets.

Word / PowerPoint đọc bằng markitdown (docmd.py, DOC_ENGINE); chưa cài hoặc lỗi thì python-docx / python-pptx bên dưới.
Excel luôn dùng openpyxl: mỗi trang tính một bảng Markdown."""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

from ...config import MAX_DOC_CHARS
from . import docmd
from .base import Adapter, Context, ExtractedDoc

OFFICE_EXTS = (".docx", ".pptx", ".xlsx")
MAX_SHEET_ROWS = 3000
MAX_SHEET_COLS = 30


class OfficeAdapter(Adapter):
    kind = "office"
    label = "Word / PowerPoint / Excel"
    file_exts = OFFICE_EXTS

    def extract(self, source: dict, ctx: Context) -> Iterator[ExtractedDoc]:
        path = ctx.raw_dir / source["file"]["stored_name"]
        yield from office_docs(path, source.get("title") or path.stem, source.get("url"), ctx)


def office_docs(path: Path, title: str, url: str | None, ctx: Context) -> Iterator[ExtractedDoc]:
    ext = path.suffix.lower()
    if ext == ".docx":
        text = docmd.docx(path, ctx.log)
        engine = "markitdown" if text else "python-docx"
        yield from split_doc("doc", title, text or docx_markdown(path), url, engine)
    elif ext == ".pptx":
        got = docmd.pptx(path, ctx.log)
        engine = "markitdown" if got else "python-pptx"
        text, n = got or pptx_markdown(path)
        ctx.log(f"PowerPoint {n} slide")
        yield from split_doc("slides", title, text, url, engine)
    elif ext == ".xlsx":
        yield from xlsx_docs(path, title, url, ctx, "openpyxl")
    else:
        raise ValueError(f"Chưa đọc được định dạng {ext}")


def split_doc(key: str, title: str, text: str, url: str | None, engine: str) -> Iterator[ExtractedDoc]:
    """Tài liệu quá dài: chia theo đoạn để mỗi phần không quá MAX_DOC_CHARS."""
    if len(text) <= MAX_DOC_CHARS:
        yield ExtractedDoc(key=key, title=title, text=text, url=url, engine=engine)
        return
    parts, buf, size = [], [], 0
    for para in text.split("\n\n"):
        if size + len(para) > MAX_DOC_CHARS and buf:
            parts.append("\n\n".join(buf))
            buf, size = [], 0
        buf.append(para)
        size += len(para) + 2
    parts.append("\n\n".join(buf))
    for i, part in enumerate(parts, 1):
        yield ExtractedDoc(key=f"{key}-{i}", title=f"{title} (phần {i}/{len(parts)})", text=part, url=url,
                           engine=engine, meta={"part": i, "parts": len(parts)})


def md_cell(v) -> str:
    if v is None:
        return ""
    return str(v).replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()


def md_table(rows: list[list]) -> str:
    rows = [r for r in rows if any(md_cell(c) for c in r)]
    if not rows:
        return ""
    # bỏ cột trống bên phải (Google Sheets / Excel hay xuất thừa cột)
    width = max(max((i + 1 for i, c in enumerate(r) if md_cell(c)), default=0) for r in rows)
    rows = [r[:width] for r in rows]
    rows = [[md_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def docx_markdown(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    out: list[str] = []
    for el in d.element.body.iterchildren():   # giữ đúng thứ tự đoạn văn / bảng
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                out.append("#" * min(int(style[-1]), 6) + " " + text)
            elif style == "Title":
                out.append("# " + text)
            elif "List" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            t = Table(el, d)
            out.append(md_table([[c.text for c in row.cells] for row in t.rows]))
    return "\n\n".join(x for x in out if x)


def pptx_markdown(path: Path) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"## Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.append(md_table([[c.text for c in row.cells] for row in shape.table.rows]))
        if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
            parts.append(f"> Ghi chú người trình bày: {notes}")
        out.append("\n\n".join(parts))
    return "\n\n".join(out), len(prs.slides)


def xlsx_docs(path: Path, title: str, url: str | None, ctx: Context, engine: str) -> Iterator[ExtractedDoc]:
    """Mỗi trang tính một tài liệu dạng bảng Markdown."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ctx.log(f"Bảng tính {len(wb.sheetnames)} trang: {', '.join(wb.sheetnames)[:200]}")
    for idx, ws in enumerate(wb.worksheets, 1):
        rows, truncated = [], False
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_SHEET_ROWS:
                truncated = True
                break
            rows.append(list(row[:MAX_SHEET_COLS]))
        table = md_table(rows)
        if not table:
            continue
        note = f"\n\n(Đã cắt ở {MAX_SHEET_ROWS} dòng đầu)" if truncated else ""
        text = f"## Trang tính: {ws.title}\n\n{table}{note}"[:MAX_DOC_CHARS]
        yield ExtractedDoc(key=f"sheet{idx}", title=f"{title} — {ws.title}", text=text, url=url, engine=engine,
                           meta={"sheet": ws.title, "rows": len(rows), "truncated": truncated})
    wb.close()
