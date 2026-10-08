"""Sinh file mẫu nhỏ (docx / pptx / pdf / xlsx) có tiêu đề + bảng để test bộ đọc tài liệu.

Không lưu file nhị phân trong repo: mỗi test tự sinh vào tmp_path. PDF viết tay (không cần reportlab):
một trang chữ lớn làm tiêu đề, một đoạn văn và một bảng kẻ ô 3 cột; trang 2 để trống (giả trang scan).
"""

from __future__ import annotations

from pathlib import Path

TABLE = [["Hang muc", "So luong", "Don gia"],
         ["Loc gio", "2", "150000"],
         ["Dau may", "4", "320000"],
         ["Ma phanh", "1", "900000"]]


def make_docx(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Quy trình bảo dưỡng", level=1)
    d.add_paragraph("Bảng vật tư cho lần bảo dưỡng 10.000 km.")
    d.add_heading("Vật tư", level=2)
    t = d.add_table(rows=len(TABLE), cols=3)
    for r, row in enumerate(TABLE):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_paragraph("Kiểm tra lốp", style="List Bullet")
    d.add_paragraph("Kiểm tra phanh", style="List Bullet")
    d.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Đào tạo kỹ thuật viên"
    s.placeholders[1].text = "Khoá cơ bản"
    s = prs.slides.add_slide(prs.slide_layouts[5])   # chỉ có tiêu đề
    s.shapes.title.text = "Bảng giá vật tư"
    tbl = s.shapes.add_table(len(TABLE), 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(TABLE):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val
    s.notes_slide.notes_text_frame.text = "Nhấn mạnh giá má phanh"
    prs.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vat tu"
    for row in TABLE:
        ws.append(row)
    wb.save(str(path))
    return path


def _pdf(streams: list[bytes]) -> bytes:
    """PDF tối giản: mỗi phần tử `streams` là nội dung một trang A4, font Helvetica (/F1) và Helvetica-Bold (/F2)."""
    n = len(streams)
    objs: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{5 + 2 * i} 0 R" for i in range(n))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    for i, s in enumerate(streams):
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {6 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(s) + s + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def _text(x: float, y: float, s: str, font: str = "F1", size: int = 11) -> str:
    return f"BT /{font} {size} Tf {x} {y} Td ({s}) Tj ET\n"


def make_pdf(path: Path, blank_page: bool = False) -> Path:
    c = _text(72, 770, "Bao duong dinh ky", "F2", 20)
    c += _text(72, 740, "Bang vat tu cho lan bao duong 10.000 km, ap dung cho xe con duoi 9 cho.")
    # bảng kẻ ô: 4 dòng x 3 cột, mỗi ô rộng 150, cao 24
    x0, y0, w, h = 72, 700, 150, 24
    rows, cols = len(TABLE), 3
    lines = "0.5 w\n"
    for r in range(rows + 1):
        lines += f"{x0} {y0 - r * h} m {x0 + cols * w} {y0 - r * h} l S\n"
    for col in range(cols + 1):
        lines += f"{x0 + col * w} {y0} m {x0 + col * w} {y0 - rows * h} l S\n"
    cells = "".join(_text(x0 + col * w + 6, y0 - r * h - 16, val, "F2" if r == 0 else "F1")
                    for r, row in enumerate(TABLE) for col, val in enumerate(row))
    c += lines + cells
    c += _text(72, 580, "Ghi chu: thay loc gio moi 10.000 km hoac khi bi ban.")
    streams = [c.encode("latin-1")]
    if blank_page:
        streams.append(b"")
    path.write_bytes(_pdf(streams))
    return path
