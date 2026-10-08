"""Xem trước file Office trong Kho tư liệu (GET /kb/sources/{id}/preview/{path}, app/kb/preview.py)."""

from __future__ import annotations

import time

import pytest

from app import db
from app.kb import pipeline, preview
from app.kb.pipeline import raw_dir, sources
from app.spaces import personal_space
from tests.conftest import make_user


@pytest.fixture
def env(tmp_path, monkeypatch, client):
    monkeypatch.setattr(pipeline, "RAW_DIR", tmp_path / "raw")   # không đụng data/raw thật
    monkeypatch.delenv("PREVIEW_DIR", raising=False)
    user = make_user("a")
    sid = sources.insert_one({"space_id": personal_space(user)["_id"], "created_by": user["_id"], "kind": "file",
                              "lane": "light", "url": None, "title": "Tệp", "status": "extracted", "priority": 0,
                              "logs": [], "options": {}, "created_at": db.now()}).inserted_id
    base = raw_dir(sid)
    base.mkdir(parents=True)
    client.login(user)
    return user, sid, base


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Doanh thu"
    ws["A1"] = "Tiêu đề gộp"
    ws.merge_cells("A1:C1")
    ws["A2"], ws["B2"], ws["C2"] = "<script>alert(1)</script>", 1234567, 0.5
    ws["C2"].number_format = "0.00%"
    ws["B2"].number_format = "#,##0"
    ws["A3"] = "=1+1"
    ws2 = wb.create_sheet("Chi phí & <khác>")
    ws2["A1"] = "Tháng"
    wb.save(path)


def url(sid, path):
    return f"/api/kb/sources/{sid}/preview/{path}"


def test_xlsx_to_html_tabs_escape_merge(env, client):
    _, sid, base = env
    (base / "files").mkdir()
    make_xlsx(base / "files" / "bảng tính.xlsx")
    r = client.get(url(sid, "files/bảng tính.xlsx"))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/html")
    assert "charset=utf-8" in r.headers["content-type"]
    h = r.text
    assert "<script>" not in h and "&lt;script&gt;alert(1)&lt;/script&gt;" in h
    assert ">Doanh thu</label>" in h and ">Chi phí &amp; &lt;khác&gt;</label>" in h
    assert 'id="t0" checked' in h and 'id="t1"' in h
    assert 'colspan="3">Tiêu đề gộp</td>' in h
    assert "1,234,567" in h and "50.00%" in h
    assert "<th>A</th><th>B</th><th>C</th>" in h
    assert "prefers-color-scheme:dark" in h
    assert "default-src 'none'" in r.headers["content-security-policy"]


def test_xlsx_truncated_notice(env, client, monkeypatch):
    import openpyxl
    _, sid, base = env
    monkeypatch.setattr(preview, "MAX_ROWS", 5)
    wb = openpyxl.Workbook()
    for i in range(1, 11):
        wb.active.cell(i, 1, i)
    wb.save(base / "dai.xlsx")
    h = client.get(url(sid, "dai.xlsx")).text
    assert "5/10 dòng" in h and "<th>6</th>" not in h


def test_csv_vietnamese(env, client):
    _, sid, base = env
    (base / "a.csv").write_bytes("Tên;Số lượng\nPhụ tùng <b>;12\n".encode("utf-8-sig"))
    (base / "b.csv").write_bytes("Tên,Giá\nĐèn,5\n".encode("cp1258"))
    h = client.get(url(sid, "a.csv")).text
    assert "Phụ tùng &lt;b&gt;" in h and '<td class="n">12</td>' in h and "<td>Số lượng</td>" in h
    assert "Đèn" in client.get(url(sid, "b.csv")).text


def test_permissions_traversal_and_unsupported(env, client):
    _, sid, base = env
    make_xlsx(base / "x.xlsx")
    (base / "ghi-chu.txt").write_text("x")
    (base.parent / "bi-mat.xlsx").write_bytes((base / "x.xlsx").read_bytes())
    assert client.get(url(sid, "../bi-mat.xlsx")).status_code == 404
    assert client.get(url(sid, "%2E%2E/bi-mat.xlsx")).status_code == 404
    assert client.get(url(sid, "khong-co.docx")).status_code == 404
    r = client.get(url(sid, "ghi-chu.txt"))
    assert r.status_code == 415 and "Chưa hỗ trợ" in r.json()["detail"]
    client.login(make_user("b"))   # người khác không xem được nguồn -> như raw_file
    assert client.get(url(sid, "x.xlsx")).status_code == client.get(f"/api/kb/sources/{sid}/raw/x.xlsx").status_code == 404


def test_cache_outside_raw_dir_and_dropped_with_source(env, client):
    _, sid, base = env
    make_xlsx(base / "x.xlsx")
    assert client.get(url(sid, "x.xlsx")).status_code == 200
    cached = list((preview.preview_root() / str(sid)).glob("*.html"))
    assert len(cached) == 1 and base not in cached[0].parents
    assert client.get(f"/api/kb/sources/{sid}").json()["raw_files"] == ["x.xlsx"]
    # file gốc đổi -> dựng lại, bỏ bản cũ
    time.sleep(0.01)
    make_xlsx(base / "x.xlsx")
    assert client.get(url(sid, "x.xlsx")).status_code == 200
    again = list((preview.preview_root() / str(sid)).glob("*.html"))
    assert len(again) == 1
    pipeline.delete_source_data(sources.find_one({"_id": sid}))
    assert not (preview.preview_root() / str(sid)).exists()


def test_no_libreoffice_503(env, client, monkeypatch):
    _, sid, base = env
    (base / "a.docx").write_bytes(b"x")
    monkeypatch.setattr(preview, "soffice_bin", lambda: None)
    assert client.get(url(sid, "a.docx")).status_code == 503


@pytest.mark.skipif(not preview.soffice_bin(), reason="không có LibreOffice")
def test_docx_to_pdf_with_cache(env, client, monkeypatch):
    import docx
    _, sid, base = env
    d = docx.Document()
    d.add_heading("Quy trình bảo dưỡng", 1)
    d.add_paragraph("Kiểm tra dầu máy, lọc gió <và> má phanh.")
    d.save(base / "quy trình.docx")
    t = time.perf_counter()
    r = client.get(url(sid, "quy trình.docx"))
    first = time.perf_counter() - t
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].startswith("inline")
    assert r.content.startswith(b"%PDF")
    calls = []
    monkeypatch.setattr(preview, "convert", lambda *a: calls.append(a))
    t = time.perf_counter()
    r2 = client.get(url(sid, "quy trình.docx"))
    print(f"docx -> pdf: lần đầu {first:.2f}s, lần cache {time.perf_counter() - t:.3f}s")
    assert r2.content == r.content and not calls
