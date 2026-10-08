#!/usr/bin/env python3
"""
TikTok -> Text: tải toàn bộ video của một/nhiều kênh TikTok, chuyển giọng nói thành chữ,
xuất ra Excel (và tùy chọn Google Sheet).

Ví dụ:
    python tiktok_to_text.py @tenkenh
    python tiktok_to_text.py @kenh1 @kenh2 --limit 50 --cookies-from-browser chrome
    python tiktok_to_text.py @tenkenh --gsheet "https://docs.google.com/spreadsheets/d/..." --gcreds service_account.json

Chạy lại lệnh cũ sẽ bỏ qua video đã xử lý (kết quả lưu trong output/cache/).
Chỉ dùng cho nghiên cứu nội bộ — không đăng lại nội dung của kênh khác.
"""

from __future__ import annotations

import argparse
import json
import platform
import random
import re
import sys
import time
import zlib
from datetime import datetime
from pathlib import Path

# Câu Whisper hay "bịa" ra khi đoạn audio chỉ có nhạc / im lặng (học từ phụ đề YouTube)
HALLUCINATION_PATTERNS = re.compile(
    r"subscribe|đăng ký kênh|ghiền mì gõ|cảm ơn các bạn đã (theo dõi|xem)|"
    r"hẹn gặp lại các bạn|like và share|bấm chuông|la la school|vietsub|"
    # nhiều thứ tiếng: dòng ghi công phụ đề trong dữ liệu huấn luyện Whisper
    r"субтитры (создавал|сделал|подготовил)|редактор субтитров|dimatorzok|amara\.org|"
    r"字幕(由|提供|志愿者|製作)|请不吝点赞|請不吝點贊|明镜与点点|ご視聴ありがとうございました|시청해 주셔서 감사합니다",
    re.IGNORECASE,
)
HAS_LETTER = re.compile(r"[^\W\d_]")


def is_no_speech(no_speech_prob: float | None, avg_logprob: float | None) -> bool:
    """Quy tắc im lặng của Whisper: nhiều khả năng không có tiếng nói và giải mã kém tự tin -> đoạn nhạc / tiếng ồn
    bị chép thành chữ vô nghĩa (vd "rally sper certains üçLuego"). Video chỉ có nhạc nhiều khi no_speech_prob = 0
    nhưng giải mã rất kém tự tin (≈ -2, lời thật ≈ -0,2): dưới -1,5 cũng bỏ."""
    logprob = avg_logprob if avg_logprob is not None else 0
    return logprob < -1.5 or ((no_speech_prob or 0) > 0.6 and logprob < -1.0)


def is_repetition_loop(text: str) -> bool:
    """Whisper kẹt vòng lặp ("that's it, that's it, …"): chữ nén được quá nhiều. Ngưỡng 2,4 như Whisper
    dùng để bỏ kết quả giải mã hỏng; câu ngắn không xét."""
    raw = text.encode("utf-8")
    return len(raw) > 60 and len(raw) / len(zlib.compress(raw)) > 2.4

EXCEL_CELL_LIMIT = 32_000   # Excel tối đa 32.767 ký tự/ô
GSHEET_CELL_LIMIT = 49_000  # Google Sheet tối đa 50.000 ký tự/ô

DEFAULT_MODELS = {
    "mlx": "mlx-community/whisper-large-v3-turbo",
    "faster": "large-v3",
    "phowhisper": "vinai/PhoWhisper-large",
}

COLUMNS = [
    ("STT", 6),
    ("Kênh", 18),
    ("Video ID", 22),
    ("Link", 45),
    ("Ngày đăng", 12),
    ("Thời lượng (giây)", 10),
    ("Lượt xem", 12),
    ("Lượt thích", 12),
    ("Bình luận", 10),
    ("Chia sẻ", 10),
    ("Caption", 50),
    ("Nội dung lời nói", 90),
    ("Trạng thái", 25),
]


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


# ---------------------------------------------------------------------------
# Tải video (yt-dlp)
# ---------------------------------------------------------------------------

def normalize_target(target: str) -> str:
    """'@kenh', 'kenh' hoặc URL đầy đủ -> URL TikTok."""
    target = target.strip()
    if target.startswith("http"):
        return target
    return f"https://www.tiktok.com/@{target.lstrip('@')}"


def ydl_base_opts(args) -> dict:
    opts = {"quiet": True, "no_warnings": True, "noprogress": True, "ignoreerrors": False}
    if args.cookies_from_browser:
        opts["cookiesfrombrowser"] = (args.cookies_from_browser,)
    if args.cookies:
        opts["cookiefile"] = args.cookies
    return opts


def list_videos(target: str, args) -> list[dict]:
    """Quét danh sách video của kênh (chưa tải)."""
    import yt_dlp

    opts = ydl_base_opts(args) | {"extract_flat": "in_playlist", "skip_download": True}
    if args.limit:
        opts["playlistend"] = args.limit
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(target, download=False)

    entries = info.get("entries")
    if entries is None:  # link của một video đơn lẻ
        entries = [info]
    videos = []
    for e in entries:
        if not e or not e.get("id"):
            continue
        url = e.get("webpage_url") or e.get("url") or ""
        if not url.startswith("http"):
            uploader = info.get("uploader") or info.get("channel") or ""
            url = f"https://www.tiktok.com/@{uploader}/video/{e['id']}"
        videos.append({"id": str(e["id"]), "url": url})
    return videos


def single_entry(info: dict) -> dict:
    """Bài Reddit / tweet... yt-dlp trả dạng playlist (id của bài) bọc video thật (id khác) -> lấy video đầu."""
    if info.get("entries") is None:
        return info
    entries = [e for e in info["entries"] if e]
    if not entries:
        raise ValueError(f"Không có video nào trong {info.get('id')}")
    return entries[0]


def download_audio(url: str, media_dir: Path, args) -> tuple[dict, Path]:
    """Tải video và trả về (metadata, đường dẫn file).

    Không dùng 'bestaudio': với TikTok, luồng audio riêng thường là bản nhạc nền gốc (sound)
    chứ không phải âm thanh của video -> mất lời nói. Lấy bản video nhỏ nhất có kèm audio.
    YouTube không còn bản gộp video+audio -> fallback cuối cùng mới lấy luồng audio riêng
    (với YouTube luồng này chính là tiếng của video), ưu tiên bitrate thấp cho nhẹ.
    """
    import yt_dlp

    opts = ydl_base_opts(args) | {
        "format": "worst[vcodec!=none][acodec!=none]/best[acodec!=none]/ba[abr<=80]/ba",
        "outtmpl": str(media_dir / "%(id)s.%(ext)s"),
        "overwrites": False,
    }
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = single_entry(ydl.extract_info(url, download=True))
        done = info.get("requested_downloads") or [{}]
        path = Path(done[0].get("filepath") or ydl.prepare_filename(info))
    if not path.exists():  # phòng trường hợp phần mở rộng thay đổi sau khi tải
        candidates = list(media_dir.glob(f"{info['id']}.*"))
        if not candidates:
            raise FileNotFoundError(f"Không tìm thấy file đã tải cho {info['id']}")
        path = candidates[0]
    return info, path


# ---------------------------------------------------------------------------
# Chuyển giọng nói thành chữ
# ---------------------------------------------------------------------------

def pick_backend(requested: str) -> str:
    if requested != "auto":
        return requested
    if sys.platform == "darwin" and platform.machine() == "arm64":
        try:
            import mlx_whisper  # noqa: F401
            return "mlx"
        except ImportError:
            pass
    return "faster"


# chữ viết không cách giữa các từ: nối câu không chèn dấu cách
NO_SPACE_LANGS = {"zh", "ja", "yue", "th", "lo", "my", "km"}


class Transcriber:
    """Bọc 3 engine: mlx-whisper (Mac M1-M4), faster-whisper (Win/Linux/GPU), PhoWhisper (VinAI)."""

    def __init__(self, backend: str, model: str | None, language: str):
        self.backend = backend
        self.model_name = model or DEFAULT_MODELS[backend]
        self.language = None if language == "auto" else language
        self.detected_language: str | None = None   # ngôn ngữ Whisper nhận ra ở lần chuyển chữ gần nhất
        self._model = None

    def _load(self):
        if self._model is not None:
            return
        log(f"Nạp model {self.model_name} ({self.backend}) — lần đầu sẽ tải model về, hơi lâu...")
        if self.backend == "faster":
            from faster_whisper import WhisperModel
            self._model = WhisperModel(self.model_name, device="auto", compute_type="auto")
        elif self.backend == "phowhisper":
            import torch
            from transformers import pipeline
            if torch.cuda.is_available():
                device = "cuda"
            elif torch.backends.mps.is_available():
                device = "mps"
            else:
                device = "cpu"
            self._model = pipeline(
                "automatic-speech-recognition",
                model=self.model_name,
                device=device,
                dtype=torch.float32 if device == "cpu" else torch.float16,
                chunk_length_s=30,
                batch_size=8,
            )
        else:  # mlx: model được nạp lười bên trong mlx_whisper.transcribe
            self._model = True

    def transcribe(self, path: Path, prompt: str | None = None,
                   language: str | None = None) -> tuple[str, list[tuple[float, float, str]]]:
        """Trả về (toàn văn, [(start, end, text), ...]). prompt: thuật ngữ làm ngữ cảnh cho Whisper;
        language: ép ngôn ngữ riêng lượt này (lượt nghe lại khi đã biết tiếng)."""
        self._load()
        lang = language or self.language
        if self.backend == "faster":
            segments, info = self._model.transcribe(
                str(path), language=lang, vad_filter=True, beam_size=5,
                condition_on_previous_text=False, initial_prompt=prompt,
            )
            segs = [(s.start, s.end, s.text.strip()) for s in segments if not is_no_speech(s.no_speech_prob, s.avg_logprob)]
            self.detected_language = info.language
        elif self.backend == "mlx":
            import mlx_whisper
            result = mlx_whisper.transcribe(
                str(path), path_or_hf_repo=self.model_name, language=lang,
                condition_on_previous_text=False, initial_prompt=prompt,
            )
            segs = [(s["start"], s["end"], s["text"].strip()) for s in result["segments"]
                    if not is_no_speech(s.get("no_speech_prob"), s.get("avg_logprob"))]
            self.detected_language = result.get("language")
        else:
            result = self._model(
                str(path), return_timestamps=True,
                generate_kwargs={"language": self.language or None, "task": "transcribe"},
            )
            segs = []
            for c in result.get("chunks", []):
                start, end = c["timestamp"]
                segs.append((start or 0.0, end if end is not None else (start or 0.0), c["text"].strip()))
            if not segs and result.get("text"):
                segs = [(0.0, 0.0, result["text"].strip())]
            self.detected_language = self.language or "vi"   # PhoWhisper chỉ nghe tiếng Việt
        # câu không có chữ nào ("?", "...", "♪") cũng là Whisper nghe nhạc / tiếng ồn
        segs = [s for s in segs if HAS_LETTER.search(s[2]) and not HALLUCINATION_PATTERNS.search(s[2])
                and not is_repetition_loop(s[2])]
        if not segs and not lang:   # chỉ có nhạc / câu bịa: ngôn ngữ Whisper đoán không có nghĩa
            self.detected_language = None
        sep = "" if self.detected_language in NO_SPACE_LANGS else " "
        text = sep.join(t for _, _, t in segs).strip()
        return text, segs


def fmt_srt_time(sec: float) -> str:
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02}:{m:02}:{s:02},{ms:03}"


def write_srt(segs: list[tuple[float, float, str]], path: Path) -> None:
    lines = []
    for i, (start, end, text) in enumerate(segs, 1):
        lines += [str(i), f"{fmt_srt_time(start)} --> {fmt_srt_time(end)}", text, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# Xuất kết quả
# ---------------------------------------------------------------------------

def format_date(info: dict) -> str:
    if info.get("timestamp"):
        return datetime.fromtimestamp(info["timestamp"]).strftime("%d/%m/%Y")
    d = info.get("upload_date")
    if d and len(d) == 8:
        return f"{d[6:8]}/{d[4:6]}/{d[0:4]}"
    return ""


def build_row(record: dict) -> dict:
    info = record.get("info", {})
    return {
        "Kênh": info.get("uploader") or info.get("channel") or record.get("channel", ""),
        "Video ID": record["id"],
        "Link": info.get("webpage_url") or record.get("url", ""),
        "Ngày đăng": format_date(info),
        "Thời lượng (giây)": info.get("duration"),
        "Lượt xem": info.get("view_count"),
        "Lượt thích": info.get("like_count"),
        "Bình luận": info.get("comment_count"),
        "Chia sẻ": info.get("repost_count"),
        "Caption": info.get("description") or info.get("title") or "",
        "Nội dung lời nói": record.get("transcript", ""),
        "Trạng thái": record.get("status", ""),
    }


def rows_as_values(rows: list[dict], cell_limit: int) -> list[list]:
    values = [[name for name, _ in COLUMNS]]
    for i, row in enumerate(rows, 1):
        line = []
        for name, _ in COLUMNS:
            v = i if name == "STT" else row.get(name)
            if isinstance(v, str) and len(v) > cell_limit:
                v = v[:cell_limit] + " …[cắt bớt, xem file .txt]"
            line.append("" if v is None else v)
        values.append(line)
    return values


def write_excel(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Videos"
    for line in rows_as_values(rows, EXCEL_CELL_LIMIT):
        ws.append(line)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    link_col = [n for n, _ in COLUMNS].index("Link") + 1
    wrap = Alignment(vertical="top", wrap_text=True)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).alignment = wrap
        link = ws.cell(row=r, column=link_col)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
    for name in ("Lượt xem", "Lượt thích", "Bình luận", "Chia sẻ"):
        c = [n for n, _ in COLUMNS].index(name) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = "#,##0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_gsheet(rows: list[dict], sheet: str, creds: str, tab: str) -> str:
    import gspread

    gc = gspread.service_account(filename=creds)
    sh = gc.open_by_url(sheet) if sheet.startswith("http") else gc.open_by_key(sheet)
    values = rows_as_values(rows, GSHEET_CELL_LIMIT)
    try:
        ws = sh.worksheet(tab)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=tab, rows=len(values) + 10, cols=len(COLUMNS))
    ws.update(range_name="A1", values=values)
    ws.freeze(rows=1)
    return sh.url


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Tải video kênh TikTok và chuyển lời nói thành văn bản (Excel / Google Sheet).",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("targets", nargs="*", help="Kênh (@tenkenh), link kênh hoặc link video. Có thể nhập nhiều.")
    p.add_argument("--file", help="File .txt chứa danh sách kênh/link, mỗi dòng một mục")
    p.add_argument("--limit", type=int, default=0, help="Chỉ lấy N video mới nhất mỗi kênh (0 = tất cả)")
    p.add_argument("--out", default="output", help="Thư mục kết quả")
    p.add_argument("--excel", help="Tên file Excel (mặc định: output/tiktok_<thời gian>.xlsx)")
    p.add_argument("--backend", choices=["auto", "mlx", "faster", "phowhisper"], default="auto",
                   help="Engine nhận dạng giọng nói. auto = mlx trên Mac M-series, còn lại faster-whisper")
    p.add_argument("--model", help="Tên model (mặc định theo backend, vd large-v3 / small / vinai/PhoWhisper-medium)")
    p.add_argument("--language", default="auto", help="'auto' để tự nhận (mặc định) hoặc mã ngôn ngữ ép cứng (vi, en, zh...)")
    p.add_argument("--cookies-from-browser", metavar="BROWSER",
                   help="Lấy cookies từ trình duyệt (chrome, safari, firefox, edge) khi TikTok chặn")
    p.add_argument("--cookies", metavar="FILE", help="File cookies.txt (định dạng Netscape)")
    p.add_argument("--sleep", type=float, default=2.0, help="Nghỉ ngẫu nhiên tối đa N giây giữa các video để tránh bị chặn")
    p.add_argument("--keep-media", action="store_true", help="Giữ lại file video/audio sau khi xử lý")
    p.add_argument("--retry-failed", action="store_true", help="Xử lý lại các video lần trước bị lỗi")
    p.add_argument("--gsheet", help="URL hoặc ID Google Sheet để ghi kết quả")
    p.add_argument("--gcreds", default="service_account.json", help="File JSON service account Google")
    p.add_argument("--gsheet-tab", default="TikTok", help="Tên tab trong Google Sheet")
    args = p.parse_args()

    if args.file:
        lines = Path(args.file).read_text(encoding="utf-8").splitlines()
        args.targets += [ln.strip() for ln in lines if ln.strip() and not ln.strip().startswith("#")]
    if not args.targets:
        p.error("Cần nhập ít nhất một kênh, ví dụ: python tiktok_to_text.py @tenkenh")
    return args


def main() -> int:
    args = parse_args()
    out = Path(args.out)
    cache_dir, srt_dir, txt_dir, media_dir = (out / d for d in ("cache", "srt", "txt", "media"))
    for d in (cache_dir, srt_dir, txt_dir, media_dir):
        d.mkdir(parents=True, exist_ok=True)

    transcriber = Transcriber(pick_backend(args.backend), args.model, args.language)
    log(f"Engine: {transcriber.backend} / {transcriber.model_name} / ngôn ngữ: {args.language}")

    # 1. Quét danh sách video
    queue: list[dict] = []
    seen: set[str] = set()
    for raw in args.targets:
        target = normalize_target(raw)
        log(f"Quét {target} ...")
        try:
            videos = list_videos(target, args)
        except Exception as e:  # noqa: BLE001
            log(f"  ✗ Không quét được: {e}")
            log("  → Thử thêm --cookies-from-browser chrome (đã đăng nhập TikTok trên Chrome)")
            continue
        log(f"  Tìm thấy {len(videos)} video")
        for v in videos:
            if v["id"] not in seen:
                seen.add(v["id"])
                queue.append(v | {"channel": raw.lstrip("@")})

    if not queue:
        log("Không có video nào để xử lý.")
        return 1

    # 2. Tải + chuyển thành chữ từng video
    records: list[dict] = []
    total = len(queue)
    for i, v in enumerate(queue, 1):
        cache_file = cache_dir / f"{v['id']}.json"
        if cache_file.exists():
            cached = json.loads(cache_file.read_text(encoding="utf-8"))
            if cached.get("status", "").startswith("OK") or not args.retry_failed:
                records.append(cached)
                log(f"[{i}/{total}] {v['id']} — đã có, bỏ qua")
                continue

        log(f"[{i}/{total}] {v['id']} — đang tải...")
        record = {"id": v["id"], "url": v["url"], "channel": v["channel"]}
        media_path = None
        try:
            info, media_path = download_audio(v["url"], media_dir, args)
            keep = ("uploader", "channel", "webpage_url", "timestamp", "upload_date", "duration",
                    "view_count", "like_count", "comment_count", "repost_count", "description", "title")
            record["info"] = {k: info.get(k) for k in keep}

            log(f"[{i}/{total}] {v['id']} — đang nhận dạng giọng nói...")
            t0 = time.time()
            text, segs = transcriber.transcribe(media_path)
            record["transcript"] = text
            record["status"] = "OK" if text else "OK (không có lời nói)"
            write_srt(segs, srt_dir / f"{v['id']}.srt")
            (txt_dir / f"{v['id']}.txt").write_text(text, encoding="utf-8")
            log(f"[{i}/{total}] {v['id']} — xong ({time.time() - t0:.0f}s, {len(text)} ký tự)")
        except Exception as e:  # noqa: BLE001
            record["status"] = f"LỖI: {str(e)[:200]}"
            log(f"[{i}/{total}] {v['id']} — ✗ {e}")
        finally:
            if media_path and media_path.exists() and not args.keep_media:
                media_path.unlink()

        cache_file.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        records.append(record)

        if args.sleep and i < total:
            time.sleep(random.uniform(args.sleep / 2, args.sleep))

    # 3. Xuất kết quả
    rows = [build_row(r) for r in records]
    excel_path = Path(args.excel) if args.excel else out / f"tiktok_{datetime.now():%Y%m%d_%H%M}.xlsx"
    write_excel(rows, excel_path)
    ok = sum(1 for r in records if r.get("status", "").startswith("OK"))
    log(f"Đã ghi Excel: {excel_path.resolve()}  ({ok}/{len(records)} video thành công)")

    if args.gsheet:
        try:
            url = write_gsheet(rows, args.gsheet, args.gcreds, args.gsheet_tab)
            log(f"Đã ghi Google Sheet: {url}")
        except Exception as e:  # noqa: BLE001
            log(f"✗ Không ghi được Google Sheet: {e}")
            log("  → Kiểm tra file service_account.json và đã share Sheet cho email của service account (quyền Editor)")

    failed = len(records) - ok
    if failed:
        log(f"Có {failed} video lỗi — chạy lại với --retry-failed để thử lại.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
